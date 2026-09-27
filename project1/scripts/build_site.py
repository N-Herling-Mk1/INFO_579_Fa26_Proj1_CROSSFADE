"""Bake the CROSSFADE project files into data/site_data.js for the docs site.

Run from the repo root after any deliverable changes:
    python scripts/build_site.py

The site itself never runs Python: every panel reads window.SITE from the
generated file, so it works on GitHub Pages and from file:// alike.
Requires: openpyxl, markdown (pip install openpyxl markdown).
"""
import json, os, sys
from datetime import date, datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "site_data.js")
STEPS = 6
def step(i, msg):
    print(f"[{i}/{STEPS}] {msg}"); sys.stdout.flush()

try:
    from openpyxl import load_workbook
    import markdown
except ImportError as e:
    sys.exit(f"missing package: {e.name}. Install with: pip install openpyxl markdown")

# ---------------------------------------------------------------- 1 tables
step(1, "reading deliverables/data.xlsx")
wb = load_workbook(os.path.join(ROOT, "deliverables", "data.xlsx"))
ws = wb["table records"]

def fmt(cell):
    v, nf = cell.value, cell.number_format
    if v is None: return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d") if "h" not in nf else v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date): return v.strftime("%Y-%m-%d")
    if isinstance(v, float) or (isinstance(v, int) and nf.startswith("0.")):
        if nf.startswith("0."): return f"{float(v):.{len(nf) - 2}f}"
        return repr(v)
    return str(v)

tables, cur, state = [], None, "title"
for row in ws.iter_rows():
    vals = [c.value for c in row]
    if all(v is None for v in vals):
        if cur: tables.append(cur); cur = None
        state = "title"; continue
    if state == "title":
        cur = {"name": str(vals[0]), "header": [], "rows": []}; state = "header"
    elif state == "header":
        cur["header"] = [str(v) for v in vals if v is not None]; state = "rows"
    else:
        cur["rows"].append([fmt(c) for c in row[:len(cur["header"])]])
if cur: tables.append(cur)
ASSOC = {"Track Score", "Genre Probability", "Station Blend", "Play"}
for t in tables: t["associative"] = t["name"] in ASSOC
entity_rows = sum(len(t["rows"]) for t in tables if not t["associative"])
print(f"      {len(tables)} tables, {sum(len(t['rows']) for t in tables)} rows "
      f"({entity_rows} in non-associative entities)")

# ---------------------------------------------------------------- 2 hero blend
step(2, "hero: Low Tide Signal genre distribution per model")
T = {t["name"]: t for t in tables}
genres = [r[1] for r in T["Genre"]["rows"]]
labels = {r[0]: r[1] for r in T["Scoring Model"]["rows"]}
track = next(r for r in T["Track"]["rows"] if r[0] == "4817")
hero = {"track_id": "4817", "title": track[2],
        "artist": next(r[1] for r in T["Artist"]["rows"] if r[0] == track[1]),
        "genres": genres, "models": []}
for tid, mid, gid, p in T["Genre Probability"]["rows"]:
    if tid != "4817": continue
    m = next((x for x in hero["models"] if x["id"] == mid), None)
    if m is None:
        ts = next(r for r in T["Track Score"]["rows"] if r[0] == tid and r[1] == mid)
        m = {"id": mid, "label": labels[mid], "probs": [0.0] * len(genres),
             "entropy": ts[4], "divergence": ts[5]}
        hero["models"].append(m)
    m["probs"][int(gid)] = float(p)
print(f"      {len(hero['models'])} model distributions for '{hero['title']}'")

# ---------------------------------------------------------------- 3 markdown
step(3, "rendering docs/DATA_DICTIONARY.md and CHANGES.md")
def md(path):
    with open(os.path.join(ROOT, path), encoding="utf-8") as f: text = f.read()
    conv = markdown.Markdown(extensions=["extra", "toc", "sane_lists"])
    html = conv.convert(text)
    toc = [{"id": t["id"], "name": t["name"]} for top in conv.toc_tokens for t in top.get("children", [])]
    return html, toc
dict_html, dict_toc = md("docs/DATA_DICTIONARY.md")
changes_html, changes_toc = md("CHANGES.md")
print(f"      dictionary {len(dict_toc)} sections, changelog {len(changes_toc)} sections")

# ---------------------------------------------------------------- 4 manifest
step(4, "file manifest")
files = []
for top in ("deliverables", "docs", "source", "scripts"):
    for dp, _, fns in os.walk(os.path.join(ROOT, top)):
        for fn in sorted(fns):
            p = os.path.join(dp, fn)
            files.append({"path": os.path.relpath(p, ROOT).replace(os.sep, "/"), "bytes": os.path.getsize(p)})
files.sort(key=lambda f: f["path"])
REQUIRED = ["requirements_analysis.pdf", "conceptual_data_model.png", "data.xlsx", "physical_data_model.png"]
status = {n: os.path.exists(os.path.join(ROOT, "deliverables", n)) for n in REQUIRED}
print(f"      {len(files)} files; deliverables present: {sum(status.values())}/4")

# ---------------------------------------------------------------- 5 checks
step(5, "sanity checks")
assert entity_rows >= 50, "fewer than 50 records in non-associative entities"
assert all(len(t["rows"]) >= 5 for t in tables if not t["associative"]), "an entity has fewer than 5 records"
for m in hero["models"]:
    assert abs(sum(m["probs"]) - 1) <= 0.001, f"model {m['id']} distribution does not sum to 1"
print("      record minimums met; hero distributions sum to 1")

# ---------------------------------------------------------------- 6 write
step(6, f"writing {os.path.relpath(OUT, ROOT)}")
site = {"tables": tables, "hero": hero, "dictionary": {"html": dict_html, "toc": dict_toc},
        "changes": {"html": changes_html, "toc": changes_toc}, "files": files, "status": status}
os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8", newline="\n") as f:
    f.write("// Generated by scripts/build_site.py. Do not edit by hand.\n")
    f.write("window.SITE = " + json.dumps(site, ensure_ascii=True, separators=(",", ":")) + ";\n")
print(f"      {os.path.getsize(OUT) / 1024:.1f} KB written")
