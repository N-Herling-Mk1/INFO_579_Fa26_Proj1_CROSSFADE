"""Build docs/DATA_DICTIONARY.md.

Types, keys and optionality come from the requirements analysis LaTeX (\\attr blocks);
extreme and typical values come from deliverables/data.xlsx; the prose comes from
scripts/dictionary_text.py. Fails if any column is missing from any of the three.
Run from project1/:
    python scripts/build_dictionary.py
then python scripts/build_site.py so the site picks it up.
"""
import os, re, sys
from collections import OrderedDict
from datetime import date, datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
TEX = os.path.join(ROOT, "source", "requirements_analysis", "requirements_analysis.tex")
XLSX = os.path.join(ROOT, "deliverables", "data.xlsx")
DOC = os.path.join(ROOT, "docs", "DATA_DICTIONARY.md")
STEPS = 5
def step(i, msg): print(f"[{i}/{STEPS}] {msg}", flush=True)

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("missing package: openpyxl. Install with: pip install openpyxl")
import dictionary_text as T

def untex(s):
    s = re.sub(r"\\q\{([^}]*)\}", r"'\1'", s)
    s = re.sub(r"\\(?:texttt|textbf|emph)\{([^}]*)\}", r"\1", s)
    s = s.replace("\\_", "_").replace("\\allowbreak", "").replace("\\ldots", "...").replace("--", "-")
    return re.sub(r"\s+", " ", s).strip()

# ------------------------------------------------------------------ 1 RA
step(1, "reading types and keys from requirements_analysis.tex")
tex = open(TEX, encoding="utf-8").read()
ARG = r"%?\s*\{((?:[^{}]|\{(?:[^{}]|\{[^{}]*\})*\})*)\}"
cols = OrderedDict(); table = None
for m in re.finditer(r"\\entityhead\{([^}]*)\}|\\attr" + ARG * 6, tex):
    if m.group(1):
        table = m.group(1); continue
    if table is None or m.group(2) == "name":
        continue
    desc = untex(m.group(7))
    fk = re.search(r"[Ff]oreign key referencing ([A-Z][a-z]+(?: [A-Z][a-z]+)?)", desc)
    cols[(table, m.group(2))] = {"type": untex(m.group(3)), "optional": m.group(4) == "Yes",
                                 "unique": m.group(5) == "Yes", "pk": m.group(6) == "Yes",
                                 "fk": fk.group(1) if fk else None}
print(f"      {len(cols)} columns in {len({t for t, _ in cols})} tables, {sum(1 for c in cols.values() if c['fk'])} foreign keys")

# ------------------------------------------------------------------ 2 sizes
step(2, "sizing every type")
INT = {"TINYINT": (1, 255), "SMALLINT": (2, 65535), "MEDIUMINT": (3, 16777215), "INT": (4, 4294967295)}
def dec_bytes(d): return (d // 9) * 4 + [0, 1, 1, 2, 2, 3, 3, 4, 4, 4][d % 9]
def size(t):
    """(bytes label, max bytes, what the type allows)"""
    m = re.match(r"(TINYINT|SMALLINT|MEDIUMINT|INT) UNSIGNED$", t)
    if m: b, hi = INT[m.group(1)]; return f"{b}", b, f"0 to {hi:,}"
    m = re.match(r"DECIMAL\((\d+),(\d+)\)$", t)
    if m:
        p, d = int(m.group(1)), int(m.group(2)); b = dec_bytes(p - d) + dec_bytes(d)
        return f"{b}", b, f"up to {('9' * (p - d) or '0')}.{'9' * d}"
    m = re.match(r"CHAR\((\d+)\)$", t)
    if m: n = int(m.group(1)); return f"{n}", n, f"exactly {n} characters"
    m = re.match(r"VARCHAR\((\d+)\)$", t)
    if m: n = int(m.group(1)); return f"length + 1 (max {n + 1})", n + 1, f"up to {n} characters"
    m = re.match(r"BINARY\((\d+)\)$", t)
    if m: n = int(m.group(1)); return f"{n}", n, f"exactly {n} bytes"
    if t.startswith("ENUM("): return "1", 1, "only the listed values"
    if t == "BIT(1)": return "1", 1, "0 or 1"
    if t == "DATE": return "3", 3, "1000-01-01 to 9999-12-31"
    if t == "DATETIME": return "8 (5 in MySQL 5.6.4+)", 8, "1000-01-01 to 9999-12-31, to the second"
    raise ValueError(f"no size rule for {t}")
for c in cols.values():
    c["bytes"], c["max"], c["allows"] = size(c["type"])

# ------------------------------------------------------------------ 3 values
step(3, "reading extreme and typical values from data.xlsx")
ws = load_workbook(XLSX, read_only=True)["table records"]
names = {t for t, _ in cols}; cur = hdr = None; data = {}; nrows = {}
for row in ws.iter_rows(values_only=True):
    vals = [v for v in row if v is not None]
    if not vals: continue
    if len(vals) == 1 and vals[0] in names: cur, hdr = vals[0], None; continue
    if cur and hdr is None: hdr = list(row); continue
    if cur:
        nrows[cur] = nrows.get(cur, 0) + 1
        for i, h in enumerate(hdr):
            if h is not None: data.setdefault((cur, h), []).append(row[i] if i < len(row) else None)
def fmt(v, t):
    if isinstance(v, datetime): return v.strftime("%Y-%m-%d") if t == "DATE" else v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date): return v.isoformat()
    if isinstance(v, (int, float)) and not isinstance(v, bool) and t.startswith("DECIMAL"):
        places = int(re.search(r",(\d+)", t).group(1)); return f"{v:.{places}f}"
    if isinstance(v, str) and v.startswith("0x") and len(v) > 20: return v[:10] + "..." + v[-4:]
    return str(v)
def code(s): return f"`{s}`"
missing = [k for k in cols if not [v for v in data.get(k, []) if v not in (None, "")]]
if missing: sys.exit(f"no sample values for {missing}")
for k, c in cols.items():
    allv = data[k]; vs = [v for v in allv if v not in (None, "")]; t = c["type"]
    empty = len(allv) - len(vs)
    uniq = list(OrderedDict.fromkeys(fmt(v, t) for v in vs))
    if t.startswith("ENUM"):
        members = re.findall(r"'([^']*)'", t)
        short = ", ".join(code(x) for x in members)
        line = f"one of {short}" + ("" if set(uniq) >= set(members) else f"; the sample uses {', '.join(code(x) for x in uniq)}")
    elif t == "BIT(1)":
        short = "0 or 1"; line = "0 (false) or 1 (true); the sample has " + (" and ".join(code(x) for x in sorted(uniq)))
    elif t.startswith(("CHAR", "BINARY")):
        short = f"always {c['max']} {'characters' if t.startswith('CHAR') else 'bytes'}, e.g. {code(uniq[0])}"
        line = f"always {c['max']} {'characters' if t.startswith('CHAR') else 'bytes'}; typical {', '.join(code(x) for x in uniq[:3])}"
    elif t.startswith("VARCHAR"):
        lo = min(vs, key=lambda v: len(str(v))); hi = max(vs, key=lambda v: len(str(v)))
        short = (f"{code(lo)} to {code(hi)}, {len(str(lo))} to {len(str(hi))} characters" if len(str(hi)) <= 20
                 else f"{len(str(lo))} to {len(str(hi))} characters")
        line = f"shortest {code(lo)} ({len(str(lo))} characters), longest {code(hi)} ({len(str(hi))}); the type allows {c['allows']}"
    else:   # numbers and dates
        lo, hi = fmt(min(vs), t), fmt(max(vs), t)
        short = f"{lo} to {hi}"
        line = f"{code(lo)} to {code(hi)} in the sample; the type allows {c['allows']}"
        mid = [x for x in uniq if x not in (lo, hi)]
        if mid: line += f"; typical {code(mid[len(mid) // 2])}"
    if empty: line += f"; empty in {empty} of {len(allv)} rows"; short += f"; {empty} empty"
    c["short"], c["line"] = short, line
    c["avg"] = (sum(len(str(v)) for v in vs) / len(vs) + 1) if t.startswith("VARCHAR") else c["max"]

# ------------------------------------------------------------------ 4 render
step(4, "rendering markdown")
prose = {(e[0], col[0]): col for e in T.ENTITIES for col in e[3]}
extra = set(prose) - set(cols); lack = set(cols) - set(prose)
if extra or lack: sys.exit(f"prose/column mismatch. no such column: {sorted(extra)}; no prose for: {sorted(lack)}")
order_ra = list(OrderedDict.fromkeys(t for t, _ in cols))
out = ["# CROSSFADE — Data dictionary", "",
       "<!-- Generated by scripts/build_dictionary.py from the RA, data.xlsx and scripts/dictionary_text.py. Edit those, not this file. -->", "",
       "## How to read this", "", T.INTRO, ""]
for name, kind, intro, colnotes in T.ENTITIES:
    if name not in order_ra: sys.exit(f"entity {name} not in RA")
    rows = [(n, cols[(name, n)]) for n, _, _ in colnotes]
    total = sum(c["max"] for _, c in rows); typ = sum(c["avg"] for _, c in rows)
    out += [f"## {name}", "",
            f"*{kind.capitalize()} · {len(rows)} columns · {nrows.get(name, 0)} sample rows · row size at most {total} bytes, about {typ:.0f} in the sample*", "",
            intro, "",
            "| Column | Type | Bytes | Extremes in data.xlsx, or typical value |", "|---|---|---|---|"]
    for n, c in rows:
        tag = " (PK)" if c["pk"] else (" (FK)" if c["fk"] else "")
        if (name, n) in T.DERIVED: tag += " *derived*"
        if (name, n) in T.SNAPSHOT: tag += " *snapshot*"
        out.append(f"| {n}{tag} | `{c['type']}` | {c['bytes']} | {c['short']} |")
    out.append("")
    for n, why, what in colnotes:
        c = cols[(name, n)]; tags = []
        if c["pk"]: tags.append("PK")
        if c["fk"]: tags.append(f"FK to {c['fk']}")
        if c["unique"] and not c["pk"]: tags.append("UNIQUE")
        if c["optional"]: tags.append("optional")
        if (name, n) in T.DERIVED: tags.append("*derived*")
        if (name, n) in T.SNAPSHOT: tags.append("*snapshot*")
        out += [f"### {n} · `{c['type']}`" + (" · " + " · ".join(tags) if tags else ""), "",
                f"- **Values:** {c['line']}.",
                f"- **Why this type:** {why}",
                f"- **What it holds:** {what}", ""]
out += [T.PATTERNS, "", T.PARSIMONY, ""]
doc = "\n".join(out)

# ------------------------------------------------------------------ 5 write
step(5, f"writing {os.path.relpath(DOC, ROOT)}")
open(DOC, "w", encoding="utf-8", newline="").write(doc)
print(f"      {len(cols)} columns, {len(T.ENTITIES)} tables, {len(doc) // 1024} KB")
