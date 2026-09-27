"""Regenerate the "Data types and parsimony" tables in docs/DATA_DICTIONARY.md.

Reads every column's declared type from the requirements analysis LaTeX (the
\\attr blocks), the declared range from its Table 8, and typical values from
deliverables/data.xlsx, then rewrites the text between the TYPES:BEGIN and
TYPES:END markers. Run from project1/:
    python scripts/build_type_glossary.py
Then run scripts/build_site.py so the site picks it up.
"""
import os, re, sys
from collections import OrderedDict
from datetime import date, datetime

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEX = os.path.join(ROOT, "source", "requirements_analysis", "requirements_analysis.tex")
XLSX = os.path.join(ROOT, "deliverables", "data.xlsx")
DOC = os.path.join(ROOT, "docs", "DATA_DICTIONARY.md")
BEGIN, END = "<!-- TYPES:BEGIN -->", "<!-- TYPES:END -->"
STEPS = 5
def step(i, msg): print(f"[{i}/{STEPS}] {msg}", flush=True)

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("missing package: openpyxl. Install with: pip install openpyxl")

def untex(s):
    s = re.sub(r"\\q\{([^}]*)\}", r"'\1'", s)
    s = re.sub(r"\\texttt\{([^}]*)\}", r"\1", s)
    s = s.replace("\\_", "_").replace("\\allowbreak", "").replace("\\ldots", "...").replace("--", "-")
    return re.sub(r"\s+", " ", s).strip()

# ------------------------------------------------------------------ 1 types
step(1, "reading declared types from requirements_analysis.tex")
tex = open(TEX, encoding="utf-8").read()
cols = OrderedDict()   # (table, column) -> dict
table = None
for m in re.finditer(r"\\entityhead\{([^}]*)\}|\\attr\{([^}]*)\}\{((?:[^{}]|\{[^{}]*\})*)\}\{([^}]*)\}\{([^}]*)\}\{([^}]*)\}", tex):
    if m.group(1):
        table = m.group(1); continue
    if table is None or m.group(2) == "name":
        continue
    cols[(table, m.group(2))] = {"type": untex(m.group(3)), "optional": m.group(4) == "Yes",
                                 "unique": m.group(5) == "Yes", "pk": m.group(6) == "Yes"}
print(f"      {len(cols)} columns in {len({t for t, _ in cols})} tables")

# Table 8 ranges for non-key attributes
s = tex.index("Table 8 --- Attribute inventory ("); e = tex.index("\\end{longtable}", s)
rng8 = {}
for line in tex[s:e].splitlines():
    m = re.match(r"\d+ & \\(?:ent|asc)\{([^}]*)\} & ([^&]+) & ([^&]+) & ([^&]+) &", line)
    if m: rng8[(m.group(1), m.group(2).strip())] = untex(m.group(4))

# ------------------------------------------------------------------ 2 bytes and ranges per type
step(2, "sizing every type")
INT = {"TINYINT": (1, 255), "SMALLINT": (2, 65535), "MEDIUMINT": (3, 16777215), "INT": (4, 4294967295)}
def dec_bytes(digits):
    return (digits // 9) * 4 + [0, 1, 1, 2, 2, 3, 3, 4, 4, 4][digits % 9]
def size(t):
    """(bytes text, max bytes, type range text)"""
    m = re.match(r"(TINYINT|SMALLINT|MEDIUMINT|INT) UNSIGNED$", t)
    if m:
        b, hi = INT[m.group(1)]; return f"{b}", b, f"0 to {hi:,}"
    m = re.match(r"DECIMAL\((\d+),(\d+)\)$", t)
    if m:
        p, d = int(m.group(1)), int(m.group(2)); b = dec_bytes(p - d) + dec_bytes(d)
        hi = ("9" * (p - d) or "0") + "." + "9" * d
        return f"{b}", b, f"-{hi} to {hi}"
    m = re.match(r"CHAR\((\d+)\)$", t)
    if m: n = int(m.group(1)); return f"{n}", n, f"exactly {n} characters"
    m = re.match(r"VARCHAR\((\d+)\)$", t)
    if m: n = int(m.group(1)); return f"length + 1 (max {n + 1})", n + 1, f"0 to {n} characters"
    m = re.match(r"BINARY\((\d+)\)$", t)
    if m: n = int(m.group(1)); return f"{n}", n, f"exactly {n} bytes"
    if t.startswith("ENUM("): k = t.count("'") // 2; return "1", 1, f"one of {k} listed values"
    if t == "BIT(1)": return "1", 1, "0 or 1"
    if t == "DATE": return "3", 3, "1000-01-01 to 9999-12-31"
    if t == "DATETIME": return "8 (slides; 5 in MySQL 5.6.4+)", 8, "1000-01-01 00:00:00 to 9999-12-31 23:59:59"
    raise ValueError(f"no size rule for {t}")
for k, c in cols.items():
    c["bytes"], c["max"], c["trange"] = size(c["type"])

# ------------------------------------------------------------------ 3 sample values
step(3, "reading typical values from data.xlsx")
ws = load_workbook(XLSX, read_only=True)["table records"]
names = {t for t, _ in cols}
cur, hdr, data = None, None, {}
for row in ws.iter_rows(values_only=True):
    vals = [v for v in row if v is not None]
    if not vals: continue
    if len(vals) == 1 and vals[0] in names: cur, hdr = vals[0], None; continue
    if cur and hdr is None: hdr = list(row); continue
    if cur:
        for i, h in enumerate(hdr):
            if h is not None and i < len(row): data.setdefault((cur, h), []).append(row[i])
def fmt(v, t):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d") if t == "DATE" else v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date): return v.isoformat()
    if isinstance(v, (int, float)) and not isinstance(v, bool) and t.startswith("DECIMAL"):
        places = int(re.search(r",(\d+)", t).group(1))
        return f"{v:.{places}f}"
    if isinstance(v, str) and v.startswith("0x") and len(v) > 20: return v[:10] + "..." + v[-4:]
    return str(v)
missing = []
for k, c in cols.items():
    vs = [v for v in data.get(k, []) if v not in (None, "")]
    if not vs: missing.append(k); c["typical"] = "(none)"; c["obs"] = ""; continue
    uniq = list(OrderedDict.fromkeys(fmt(v, c["type"]) for v in vs))
    pick = uniq if len(uniq) <= 3 else [uniq[0], uniq[len(uniq) // 2], uniq[-1]]
    c["typical"] = ", ".join(f"`{p}`" for p in pick)
    nums = [v for v in vs if isinstance(v, (int, float)) and not isinstance(v, bool)]
    if nums and len(nums) == len(vs) and not c["type"].startswith("BIT"):
        c["obs"] = f"{fmt(min(nums), c['type'])} to {fmt(max(nums), c['type'])}"
    elif all(isinstance(v, (date, datetime)) for v in vs):
        c["obs"] = f"{fmt(min(vs), c['type'])} to {fmt(max(vs), c['type'])}"
    elif c["type"].startswith(("VARCHAR", "CHAR")):
        L = [len(str(v)) for v in vs]
        c["obs"] = f"always {L[0]} characters" if min(L) == max(L) else f"{min(L)} to {max(L)} characters"
    else:
        c["obs"] = ""
    c["avg"] = sum(len(str(v)) for v in vs) / len(vs) + 1 if c["type"].startswith("VARCHAR") else c["max"]
if missing: sys.exit(f"no sample values for {missing}")
print(f"      sample values found for all {len(cols)} columns")

# design ranges for key columns (Table 8 lists non-key attributes only)
KEYRANGE = {"Genre ID": "0 to 9, one per genre (the models' class index)",
            "Model ID": "1 upward, in deployment order (up to 255)",
            "Artist ID": "generated, up to 16,777,215", "Track ID": "generated, up to 4,294,967,295",
            "Listener ID": "generated, up to 4,294,967,295", "Station ID": "generated, up to 4,294,967,295",
            "Played At": "date and time the stream began"}

# ------------------------------------------------------------------ 4 render
step(4, "rendering markdown tables")
TYPEREF = [("TINYINT UNSIGNED", "1", "0 to 255"), ("SMALLINT UNSIGNED", "2", "0 to 65,535"),
           ("MEDIUMINT UNSIGNED", "3", "0 to 16,777,215"), ("INT UNSIGNED", "4", "0 to 4,294,967,295"),
           ("DECIMAL(M,D)", "digits packed 9 per 4 bytes", "M digits, D after the point"),
           ("CHAR(n)", "n", "exactly n characters"), ("VARCHAR(n)", "length + 1", "0 to n characters"),
           ("BINARY(n)", "n", "exactly n bytes"), ("ENUM(...)", "1", "one listed value (up to 255 values)"),
           ("BIT(1)", "1", "0 or 1"), ("DATE", "3", "1000-01-01 to 9999-12-31"),
           ("DATETIME", "8 (slides; 5 in MySQL 5.6.4+)", "years 1000 to 9999"),
           ("TIMESTAMP (not used)", "4", "1970-01-01 to 2038-01-19")]
used = {}
for (t, n), c in cols.items():
    fam = re.sub(r"\(\d+(,\d+)?\)", lambda m: "(M,D)" if "," in m.group(0) else "(n)", c["type"])
    fam = "ENUM(...)" if fam.startswith("ENUM") else ("BIT(1)" if c["type"] == "BIT(1)" else fam)
    used.setdefault(fam, []).append(f"{t}.{n}")
out = ["### Type reference", "",
       "Byte counts are the Unit 4 figures; DECIMAL, BIT and ENUM follow the MySQL manual. String sizes assume one byte per character.", "",
       "| Type | Bytes | Range | Columns using it |", "|---|---|---|---|"]
for ty, b, r in TYPEREF:
    out.append(f"| {ty} | {b} | {r} | {len(used.get(ty, []))} |")
out += ["", "### Every column", "",
        "Declared range is what the design allows; typical values are taken from `data.xlsx`.", ""]
for tname in OrderedDict.fromkeys(t for t, _ in cols):
    rows = [(n, c) for (t, n), c in cols.items() if t == tname]
    fixed = sum(c["max"] for _, c in rows); typ = sum(c["avg"] for _, c in rows)
    out += [f"#### {tname}", "",
            f"Row size: {fixed} bytes at most, about {typ:.0f} bytes for the sample rows.", "",
            "| Column | Type | Bytes | Declared range | Typical values |", "|---|---|---|---|---|"]
    for n, c in rows:
        role = " (PK)" if c["pk"] else ""
        rng = rng8.get((tname, n)) or KEYRANGE.get(n, c["trange"])
        if c["obs"]: rng += f"; sample {c['obs']}"
        out.append(f"| {n}{role} | {c['type']} | {c['bytes']} | {rng} | {c['typical']} |")
    out.append("")
block = "\n".join(out).rstrip() + "\n"

# ------------------------------------------------------------------ 5 write
step(5, f"writing {os.path.relpath(DOC, ROOT)}")
doc = open(DOC, encoding="utf-8").read()
if BEGIN not in doc or END not in doc:
    sys.exit(f"markers {BEGIN} / {END} not found in {DOC}")
doc = doc[:doc.index(BEGIN) + len(BEGIN)] + "\n" + block + doc[doc.index(END):]
open(DOC, "w", encoding="utf-8", newline="").write(doc)
print(f"      {len(cols)} columns, {len(OrderedDict.fromkeys(t for t, _ in cols))} tables written")
