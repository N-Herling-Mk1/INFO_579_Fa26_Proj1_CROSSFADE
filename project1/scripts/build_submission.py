"""Build the Project 1 submission zip exactly as the brief specifies, and check it.

INFO 579 Project 1 brief: "Compress all the deliverables into a single zip file ... The name
file must follow the format: project01_groupcode.zip". The four deliverables and their
required names come from the same brief. Run from project1/:
    python scripts/build_submission.py
Writes submission/project01_<group code>.zip and exits 1 if any check fails.
"""
import os, sys, zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GROUP_CODE = "info579_group_1"          # D2L group name
ZIP_NAME = f"project01_{GROUP_CODE}.zip"
DELIVERABLES = ["requirements_analysis.pdf", "data.xlsx", "conceptual_data_model.png", "physical_data_model.png"]
SIGNATURES = {".pdf": b"%PDF-", ".png": b"\x89PNG\r\n\x1a\n", ".xlsx": b"PK\x03\x04"}
STEPS = 4
def step(i, msg): print(f"[{i}/{STEPS}] {msg}", flush=True)

step(1, "checking the four deliverables")
errors = []
for name in DELIVERABLES:
    path = os.path.join(ROOT, "deliverables", name)
    if not os.path.exists(path):
        errors.append(f"missing: deliverables/{name}"); continue
    with open(path, "rb") as f: head = f.read(8)
    sig = SIGNATURES[os.path.splitext(name)[1]]
    if not head.startswith(sig): errors.append(f"{name} is not a real {os.path.splitext(name)[1]} file")
    print(f"      ok  {name:28} {os.path.getsize(path) / 1024:8.1f} KB")
if errors: sys.exit("      " + "\n      ".join(errors))

step(2, "checking data.xlsx follows data_template.xlsx (one sheet, 'table records')")
from openpyxl import load_workbook
wb = load_workbook(os.path.join(ROOT, "deliverables", "data.xlsx"), read_only=True)
if wb.sheetnames != ["table records"]: sys.exit(f"      data.xlsx sheets are {wb.sheetnames}")
print("      ok  single sheet 'table records'")

step(3, f"writing submission/{ZIP_NAME}")
out_dir = os.path.join(ROOT, "submission"); os.makedirs(out_dir, exist_ok=True)
out = os.path.join(out_dir, ZIP_NAME)
with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
    for name in DELIVERABLES:                      # files at the zip's top level, no folders
        z.write(os.path.join(ROOT, "deliverables", name), arcname=name)

step(4, "re-reading the zip")
with zipfile.ZipFile(out) as z:
    names = z.namelist(); bad = z.testzip()
if sorted(names) != sorted(DELIVERABLES) or bad:
    sys.exit(f"      zip contents wrong: {names} (corrupt member: {bad})")
print(f"      ok  {ZIP_NAME}: {len(names)} files, {os.path.getsize(out) / 1024:.1f} KB, nothing else inside")
