"""Cross-check every column across the requirements analysis, data.xlsx and the data dictionary.

The RA entity tables (the \\attr blocks) are the reference. Every other place a column
appears is compared with them. Run from project1/:
    python scripts/check_consistency.py
Exits 1 if any check fails.
"""
import os, re, sys
from collections import OrderedDict, Counter
from datetime import date, datetime
from decimal import Decimal

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEX = os.path.join(ROOT, "source", "requirements_analysis", "requirements_analysis.tex")
XLSX = os.path.join(ROOT, "deliverables", "data.xlsx")
DOC = os.path.join(ROOT, "docs", "DATA_DICTIONARY.md")
from openpyxl import load_workbook

CHECKS = []
def check(name):
    def wrap(fn): CHECKS.append((name, fn)); return fn
    return wrap

def untex(s):
    s = re.sub(r"\\q\{([^}]*)\}", r"'\1'", s)
    s = re.sub(r"\\(?:ent|asc|rae|rel|textbf|emph|texttt)\{([^}]*)\}", r"\1", s)
    s = s.replace("\\_", "_").replace("{\\small (associative)}", "").replace("\\allowbreak", "")
    return re.sub(r"\s+", " ", s).strip()

tex = open(TEX, encoding="utf-8").read()
body = tex[tex.index("\\begin{document}"):]
ARG = r"%?\s*\{((?:[^{}]|\{(?:[^{}]|\{[^{}]*\})*\})*)\}"
RA = OrderedDict(); cur = None
for m in re.finditer(r"\\entityhead\{([^}]*)\}|\\attr" + ARG * 6, body):
    if m.group(1): cur = m.group(1); RA[cur] = OrderedDict(); continue
    desc = untex(m.group(7))
    fk = re.search(r"[Ff]oreign key referencing ([A-Z][a-z]+(?: [A-Z][a-z]+)?)", desc)
    RA[cur][m.group(2)] = dict(type=untex(m.group(3)), opt=m.group(4) == "Yes", uniq=m.group(5) == "Yes",
                               pk=m.group(6) == "Yes", fk=fk.group(1) if fk else None)
ALL = [(t, c, d) for t, cols in RA.items() for c, d in cols.items()]
NONKEY = [(t, c, d) for t, c, d in ALL if not d["pk"] and not d["fk"]]

def table_rows(title):
    s = body.index(title); e = body.index("\\end{longtable}", s)
    return [untex(l) for l in body[s:e].splitlines() if re.match(r"\s*\\?(?:rel\{)?[A-Z]?\d+ &|\s*\\(?:ent|asc)\{", l)]

@check("RA entity tables: 10 tables, 59 columns, 38 non-key")
def _():
    n = (len(RA), len(ALL), len(NONKEY))
    return [] if n == (10, 59, 38) else [f"got tables/columns/non-key = {n}"]

@check("RA Table 1 and Table 3: attribute counts and primary keys")
def _():
    errs = []
    for title in ("Table 1 --- Entities", "Table 3 --- Associative entities"):
        for r in table_rows(title):
            f = [x.strip() for x in r.split("&")]
            t, n, pk = f[1], int(f[2]), f[3]
            want = sum(1 for c, d in RA[t].items() if not d["pk"] and not d["fk"])
            pks = " + ".join(c for c, d in RA[t].items() if d["pk"])
            if n != want: errs.append(f"{title[:7]} {t}: attributes {n}, should be {want}")
            if pk != pks: errs.append(f"{title[:7]} {t}: key '{pk}', entity table says '{pks}'")
    return errs

@check("RA entity summary (page 2): column lists and non-key counts")
def _():
    errs = []
    s = body.index("\\subsection{Entities}"); e = body.index("\\end{longtable}", s)
    for line in body[s:e].splitlines():
        m = re.match(r"\d+ & (.+?) & (\d+) & (.+?) \\\\", line)
        if not m: continue
        t = untex(m.group(1)); listed = [re.sub(r" \(.*\)$", "", x.strip()) for x in re.split(r",(?![^()]*\))", untex(m.group(3)))]
        if listed != list(RA[t]): errs.append(f"{t}: summary lists {listed}, entity table has {list(RA[t])}")
        want = sum(1 for d in RA[t].values() if not d["pk"] and not d["fk"])
        if int(m.group(2)) != want: errs.append(f"{t}: summary count {m.group(2)}, should be {want}")
    head = re.search(r"Non-key attributes \(n=(\d+)\)", body)
    if not head or int(head.group(1)) != len(NONKEY): errs.append(f"summary header n={head and head.group(1)}")
    return errs

@check("RA Table 6: foreign keys match the entity tables")
def _():
    listed = set()
    for r in table_rows("Table 6 --- Foreign keys"):
        f = [x.strip() for x in r.split("&")]
        for col in f[1].split(" + "): listed.add((f[0], col, f[2]))
    actual = {(t, c, d["fk"]) for t, c, d in ALL if d["fk"]}
    errs = [f"in Table 6, not in entity tables: {x}" for x in sorted(listed - actual)]
    errs += [f"in entity tables, not in Table 6: {x}" for x in sorted(actual - listed)]
    if len(actual) != 14: errs.append(f"{len(actual)} foreign key columns, expected 14")
    return errs

@check("RA Table 8: every non-key attribute, same type, numbered 1..38")
def _():
    errs = []; seen = []
    for r in table_rows("Table 8 --- Attribute inventory ("):
        f = [x.strip() for x in r.split("&")]
        num, t, c, ty = int(f[0]), f[1], f[2], f[3]
        seen.append((t, c)); d = RA.get(t, {}).get(c)
        if d is None: errs.append(f"#{num} {t}.{c} not in entity tables"); continue
        if d["type"] != ty: errs.append(f"#{num} {t}.{c}: Table 8 {ty}, entity table {d['type']}")
        if num != len(seen): errs.append(f"#{num} out of sequence")
    missing = [(t, c) for t, c, _ in NONKEY if (t, c) not in seen]
    if missing: errs.append(f"missing from Table 8: {missing}")
    return errs

@check("RA Table 9: type families and column counts")
def _():
    fam = Counter()
    for _, _, d in ALL:
        ty = d["type"]
        fam["Integer" if "INT" in ty else "Fixed-point" if ty.startswith("DECIMAL") else
            "String" if "CHAR" in ty else "Binary" if ty.startswith("BINARY") else
            "Date and time" if ty.startswith("DATE") else "Enumeration" if ty.startswith("ENUM") else "Boolean"] += 1
    errs = []
    for r in table_rows("Table 9 --- Data types used"):
        pass
    s = body.index("Table 9 --- Data types used"); e = body.index("\\end{longtable}", s)
    for line in body[s:e].splitlines():
        m = re.match(r"(Integer|Fixed-point|String|Binary|Date and time|Enumeration|Boolean) & (.+?) & (\d+) &", line)
        if m and int(m.group(3)) != fam[m.group(1)]:
            errs.append(f"{m.group(1)}: Table 9 says {m.group(3)}, actual {fam[m.group(1)]}")
        if m and m.group(1) == "String":
            listed = set(re.findall(r"(?:VAR)?CHAR\(\d+\)", m.group(2)))
            actual = {d["type"] for _, _, d in ALL if "CHAR" in d["type"]}
            if listed != actual: errs.append(f"String types listed {sorted(listed)}, used {sorted(actual)}")
    return errs

# ---------------------------------------------------------------- data.xlsx
ws = load_workbook(XLSX, read_only=True)["table records"]
X = OrderedDict(); cur = hdr = None
for row in ws.iter_rows(values_only=True):
    vals = [v for v in row if v is not None]
    if not vals: continue
    if len(vals) == 1 and vals[0] in RA: cur, hdr = vals[0], None; continue
    if cur and hdr is None: hdr = [h for h in row if h is not None]; X[cur] = {"hdr": hdr, "rows": []}; continue
    if cur: X[cur]["rows"].append(list(row[:len(hdr)]) + [None] * (len(hdr) - len(row)))

@check("data.xlsx: same tables, same columns in the same order")
def _():
    errs = []
    if set(X) != set(RA): errs.append(f"tables differ: {set(X) ^ set(RA)}")
    for t in RA:
        if t in X and X[t]["hdr"] != list(RA[t]): errs.append(f"{t}: xlsx {X[t]['hdr']} vs RA {list(RA[t])}")
    return errs

def fits(v, ty):
    if isinstance(v, bool): v = int(v)
    m = re.match(r"(TINYINT|SMALLINT|MEDIUMINT|INT) UNSIGNED$", ty)
    if m: return isinstance(v, int) and 0 <= v <= {"TINYINT": 255, "SMALLINT": 65535, "MEDIUMINT": 16777215, "INT": 4294967295}[m.group(1)]
    m = re.match(r"DECIMAL\((\d+),(\d+)\)$", ty)
    if m:
        p, s = int(m.group(1)), int(m.group(2)); d = Decimal(repr(v)) if isinstance(v, float) else Decimal(v)
        return -d.as_tuple().exponent <= s and abs(d) < Decimal(10) ** (p - s)
    m = re.match(r"(VAR)?CHAR\((\d+)\)$", ty)
    if m: return isinstance(v, str) and (len(v) <= int(m.group(2)) if m.group(1) else len(v) == int(m.group(2)))
    m = re.match(r"BINARY\((\d+)\)$", ty)
    if m: return isinstance(v, str) and re.fullmatch(r"0x[0-9A-Fa-f]{%d}" % (2 * int(m.group(1))), v) is not None
    if ty.startswith("ENUM"): return v in re.findall(r"'([^']*)'", ty)
    if ty == "BIT(1)": return v in (0, 1)
    if ty == "DATE": return isinstance(v, (date, datetime)) and (not isinstance(v, datetime) or v.time() == datetime.min.time())
    if ty == "DATETIME": return isinstance(v, (date, datetime))
    return False

@check("data.xlsx: every value fits its declared type; empty only where optional")
def _():
    errs = []
    for t, cols in RA.items():
        for i, (c, d) in enumerate(cols.items()):
            for r in X[t]["rows"]:
                v = r[i]
                if v in (None, ""):
                    if not d["opt"]: errs.append(f"{t}.{c}: empty but not optional")
                elif not fits(v, d["type"]): errs.append(f"{t}.{c}: {v!r} does not fit {d['type']}")
    return errs[:15]

@check("data.xlsx: primary keys and UNIQUE columns have no duplicates")
def _():
    errs = []
    for t, cols in RA.items():
        names = list(cols); rows = X[t]["rows"]
        pk = [names.index(c) for c, d in cols.items() if d["pk"]]
        keys = [tuple(r[i] for i in pk) for r in rows]
        dup = [k for k, n in Counter(keys).items() if n > 1]
        if dup: errs.append(f"{t}: duplicate keys {dup[:3]}")
        for c, d in cols.items():
            if d["uniq"] and not d["pk"]:
                vals = [r[names.index(c)] for r in rows if r[names.index(c)] is not None]
                if len(vals) != len(set(vals)): errs.append(f"{t}.{c}: duplicate values")
    return errs

@check("data.xlsx: every foreign key value exists in the table it references")
def _():
    errs = []
    for t, cols in RA.items():
        names = list(cols)
        comp = [c for c, d in cols.items() if d["fk"] == "Track Score"]
        for c, d in cols.items():
            if not d["fk"] or c in comp: continue
            ref = RA[d["fk"]]; ref_pk = [x for x, dd in ref.items() if dd["pk"]][0]
            have = {r[list(ref).index(ref_pk)] for r in X[d["fk"]]["rows"]}
            bad = {r[names.index(c)] for r in X[t]["rows"] if r[names.index(c)] is not None} - have
            if bad: errs.append(f"{t}.{c} -> {d['fk']}: missing {sorted(bad)[:5]}")
        if comp:
            have = {(r[0], r[1]) for r in X["Track Score"]["rows"]}
            bad = {tuple(r[names.index(c)] for c in comp) for r in X[t]["rows"]} - have
            if bad: errs.append(f"{t} -> Track Score: missing {sorted(bad)[:5]}")
    return errs

# ---------------------------------------------------------------- dictionary
doc = open(DOC, encoding="utf-8").read()
@check("Data dictionary: every column, same order, type and key role")
def _():
    errs = []; sections = re.split(r"\n## ", doc)
    for t, cols in RA.items():
        sec = next((s for s in sections if s.startswith(t + "\n")), None)
        if sec is None: errs.append(f"no section for {t}"); continue
        heads = re.findall(r"\n### (.+?) · `([^`]+)`(.*)", sec)
        if [h[0] for h in heads] != list(cols): errs.append(f"{t}: dictionary order {[h[0] for h in heads]}")
        for name, ty, tags in heads:
            d = cols.get(name)
            if not d: continue
            if ty != d["type"]: errs.append(f"{t}.{name}: dictionary {ty}, RA {d['type']}")
            if ("PK" in tags) != d["pk"]: errs.append(f"{t}.{name}: PK tag mismatch")
            if (f"FK to {d['fk']}" in tags) != bool(d["fk"]): errs.append(f"{t}.{name}: FK tag mismatch")
            if ("optional" in tags) != d["opt"]: errs.append(f"{t}.{name}: optional tag mismatch")
        rows = re.findall(r"\n\| ([^|]+?)(?: \((?:PK|FK)\))?(?: \*\w+\*)* \| `([^`]+)` \|", sec)
        if [(r[0], r[1]) for r in rows] != [(c, d["type"]) for c, d in cols.items()]:
            errs.append(f"{t}: summary table rows differ")
    if "59 columns" not in doc: errs.append("intro does not say 59 columns")
    return errs

# ---------------------------------------------------------------- run
failed = 0
for i, (name, fn) in enumerate(CHECKS, 1):
    errs = fn()
    print(f"[{i}/{len(CHECKS)}] {'PASS' if not errs else 'FAIL'}  {name}", flush=True)
    for e in errs: print(f"        - {e}")
    failed += bool(errs)
print(f"\n{len(CHECKS) - failed}/{len(CHECKS)} checks passed")
sys.exit(1 if failed else 0)
