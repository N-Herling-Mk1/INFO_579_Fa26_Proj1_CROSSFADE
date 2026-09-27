"""CROSSFADE sample data -> data.xlsx (format of data_template.xlsx).

Derived Track Score values (Genre ID, Top Probability, Posterior Entropy,
Model Divergence) are COMPUTED from the stored Genre Probability rows, per
integrity rules I2/I3 and the mk5 divergence rule. Every rule is re-checked
before the workbook is written.
"""
import hashlib, math, random, sys
from datetime import date, datetime, timedelta

SEED = 579
rng = random.Random(SEED)
STEPS = 7
def step(i, msg):
    print(f"[{i}/{STEPS}] {msg}"); sys.stdout.flush()

# --------------------------------------------------------------------------
step(1, "fixed reference tables (Genre, Scoring Model)")
GENRES = [
    (0, "blues", "Twelve-bar forms, blue notes and call-and-response phrasing"),
    (1, "classical", "Orchestral and chamber writing in the Western art-music tradition"),
    (2, "country", "Narrative songwriting with steel guitar, fiddle and twang"),
    (3, "disco", "Four-on-the-floor dance grooves with lush strings and bass"),
    (4, "hiphop", "Rhythmic vocal delivery over sampled or programmed beats"),
    (5, "jazz", "Swing, improvisation and extended harmony"),
    (6, "metal", "Distorted guitars, heavy drums and aggressive dynamics"),
    (7, "pop", "Hook-driven songs built for broad appeal"),
    (8, "reggae", "Offbeat skank rhythm and prominent bass lines"),
    (9, "rock", "Guitar-led songs with a driving backbeat"),
]
G = {g[1]: g[0] for g in GENRES}
NG = len(GENRES)

# Real FORGE bundle properties (arch.json / metrics.json). Holdout = test accuracy.
# cfg12 fusion type is not recorded in the notes: assumed concat.
MODELS = [
    (1, "beardown",           "gated",  512, 30, "GTZAN", 0.7000, date(2026, 1, 12)),
    (2, "beardown_rrm_cfg12", "concat", 384, 30, "GTZAN", 0.6800, date(2026, 2, 16)),
    (3, "beardown_rrm_cfg17", "concat", 768, 30, "GTZAN", 0.7600, date(2026, 3, 9)),
    (4, "beardown_rrm",       "concat", 768, 30, "GTZAN", 0.7600, date(2026, 3, 23)),
    (5, "beardown_3sec",      "concat", 384,  3, "GTZAN", 0.7887, date(2026, 6, 23)),
]
M = {m[0]: m for m in MODELS}

# --------------------------------------------------------------------------
step(2, "artists, tracks, listeners")
ARTISTS = [  # id, stage name, country, joined, payout rate, verified
    (1042, "Velvet Static",     "US", date(2026, 3, 14), 0.00400, 1),
    (1043, "Mara Quell",        "GB", date(2026, 1, 20), 0.00350, 1),
    (1044, "Low Orbit Choir",   "CA", date(2026, 2, 3),  0.00380, 1),
    (1045, "Tin Harbour",       "AU", date(2026, 2, 27), 0.00320, 0),
    (1046, "Ochre Parade",      "MX", date(2026, 4, 11), 0.00360, 1),
    (1047, "Juno Fleet",        "DE", date(2026, 5, 2),  0.00300, 1),
    (1048, "Saltmarsh Tapes",   None, date(2026, 5, 19), 0.00300, 0),
    (1049, "Kestrel Ave",       "JM", date(2026, 8, 28), 0.00340, 0),   # no tracks yet
]
A = {a[0]: a for a in ARTISTS}

# id, artist, title, duration, release, status, blend(genre->weight), models scoring it
TRACKS = [
    (4817, 1042, "Low Tide Signal",       214, date(2026, 4, 2),  "scored",  {"jazz": .58, "hiphop": .30}, [1, 2, 4, 5]),
    (4818, 1042, "Copper Lanterns",       187, date(2026, 4, 20), "scored",  {"jazz": .45, "blues": .40}, [1, 4]),
    (4819, 1043, "Glass Meridian",        241, date(2026, 2, 25), "scored",  {"classical": .55, "pop": .30}, [1, 2]),
    (4820, 1043, "Second Weather",        199, date(2026, 3, 30), "scored",  {"pop": .50, "disco": .35}, [2, 4]),
    (4821, 1044, "Cathedral Static",      312, date(2026, 3, 1),  "scored",  {"classical": .50, "metal": .35}, [1, 2]),
    (4822, 1044, "Hymn for a Server Room",268, date(2026, 5, 15), "scored",  {"classical": .70, "jazz": .15}, [4, 5]),
    (4823, 1045, "Dust on the Dial",      205, date(2026, 3, 18), "scored",  {"country": .55, "rock": .30}, [2, 4]),
    (4824, 1045, "Harbour Lights Waltz",  176, date(2026, 5, 6),  "scored",  {"country": .75, "blues": .10}, [4]),
    (4825, 1046, "Marigold Dub",          233, date(2026, 5, 1),  "scored",  {"reggae": .55, "hiphop": .30}, [1, 4]),
    (4826, 1046, "Plaza at Noon",         198, date(2026, 6, 30), "scored",  {"disco": .45, "reggae": .35}, [4, 5]),
    (4827, 1047, "Stahlwerk",             254, date(2026, 5, 25), "scored",  {"metal": .60, "rock": .30}, [1, 4]),
    (4828, 1047, "Afterimage Club",       221, date(2026, 7, 10), "scored",  {"disco": .50, "pop": .30}, [2, 5]),
    (4829, 1048, "Tapeworm Blues",        189, date(2026, 6, 5),  "scored",  {"blues": .50, "rock": .35}, [2, 4]),
    (4830, 1048, "Brine",                 163, date(2026, 9, 21), "awaiting scoring", None, []),
    (4831, 1048, "Untitled Loop 7",        48, date(2026, 9, 22), "rejected", None, []),
]
T = {t[0]: t for t in TRACKS}

LISTENERS = [  # id, email, display, tier, signup
    (20931, "j.rivera@example.com",   "jrivera",   "standard", date(2026, 5, 9)),
    (20932, "amara.o@example.com",    "amara_o",   "premium",  date(2026, 4, 2)),
    (20933, "k.lindqvist@example.com","klind",     "free",     date(2026, 4, 18)),
    (20934, "priya.m@example.com",    "priyam",    "premium",  date(2026, 5, 27)),
    (20935, "tomas.b@example.com",    "tbarros",   "standard", date(2026, 6, 3)),
    (20936, "wen.zhao@example.com",   "wenz",      "free",     date(2026, 6, 20)),
    (20937, "d.okafor@example.com",   "dokafor",   "standard", date(2026, 7, 8)),
    (20938, "h.ito@example.com",      "hito",      "free",     date(2026, 8, 15)),  # no stations
]
L = {l[0]: l for l in LISTENERS}

# --------------------------------------------------------------------------
step(3, "genre distributions -> Genre Probability (stored in 1e-4 units)")
U = 10000
def to_units(p):
    """Round a distribution to 4 dp so it sums to exactly 1.0000; no zero cells."""
    raw = [max(x, 0.0) for x in p]; s = sum(raw); raw = [x / s for x in raw]
    u = [max(1, round(x * U)) for x in raw]
    d = U - sum(u)
    order = sorted(range(NG), key=lambda i: -u[i])
    i = 0
    while d != 0:
        k = order[i % NG]
        if d > 0: u[k] += 1; d -= 1
        elif u[k] > 1: u[k] -= 1; d += 1
        i += 1
    return u

def ent(u):   # normalised Shannon entropy of a stored distribution
    return -sum((x / U) * math.log2(x / U) for x in u if x) / math.log2(NG)

def jsd(u, v):  # base-2 Jensen-Shannon divergence, in [0,1]
    def kl(a, b): return sum((x / U) * math.log2(x / y) for x, y in zip(a, b) if x)
    m = [(x + y) / 2 for x, y in zip(u, v)]
    return 0.5 * kl(u, m) + 0.5 * kl(v, m)

def dirichlet(alpha):
    g = [rng.gammavariate(a, 1.0) for a in alpha]; s = sum(g); return [x / s for x in g]

def model_dist(blend, model_id):
    conc = {1: 60, 2: 45, 3: 80, 4: 80, 5: 110}[model_id]   # sharper = more confident model
    rest = 1 - sum(blend.values())
    base = [blend.get(GENRES[i][1], rest / (NG - len(blend))) for i in range(NG)]
    return dirichlet([conc * b + 0.05 for b in base])

DIST = {}  # (track, model) -> units
for tid, _, _, _, _, st, blend, mids in TRACKS:
    for m in mids:
        DIST[(tid, m)] = to_units(model_dist(blend, m))

# Pin the requirements-analysis example: track 4817 x beardown_3sec has
# top probability 0.5800 (jazz), entropy 0.4619 and mean divergence 0.0198.
def build_anchor():
    best = None
    for h in range(2200, 3600):                      # hiphop share, 1e-4 units
        for shape in (0.30, 0.45, 0.60, 0.75):
            rem = U - 5800 - h
            w = [shape ** k for k in range(8)]; sw = sum(w)
            others = [i for i in range(NG) if i not in (G["jazz"], G["hiphop"])]
            u = [0] * NG; u[G["jazz"]] = 5800; u[G["hiphop"]] = h
            alloc = [max(1, round(rem * x / sw)) for x in w]
            alloc[0] += rem - sum(alloc)
            if alloc[0] >= h or min(alloc) < 1: continue
            for i, a in zip(others, alloc): u[i] = a
            e = round(ent(u), 4)
            if e == 0.4619: return u
            if best is None or abs(ent(u) - 0.4619) < abs(ent(best) - 0.4619): best = u
    return best
anchor = build_anchor()
assert round(ent(anchor), 4) == 0.4619 and max(anchor) == 5800
DIST[(4817, 5)] = anchor

def mix(p, q, t):
    return [(1 - t) * a / U + t * b for a, b in zip(p, q)]
dirs = {m: model_dist(T[4817][6], m) for m in (1, 2, 4)}
def div5(t):
    for m in (1, 2, 4): DIST[(4817, m)] = to_units(mix(anchor, dirs[m], t))
    return sum(jsd(anchor, DIST[(4817, m)]) for m in (1, 2, 4)) / 3
lo, hi = 0.0, 1.0
for _ in range(60):
    mid = (lo + hi) / 2
    if div5(mid) < 0.0198: lo = mid
    else: hi = mid
t_star = min((lo + k * 1e-5 for k in range(-400, 400)), key=lambda t: abs(div5(t) - 0.0198))
div5(t_star)
# the earlier models must still rank jazz first for the example to read sensibly
assert all(DIST[(4817, m)].index(max(DIST[(4817, m)])) == G["jazz"] for m in (1, 2, 4))

# --------------------------------------------------------------------------
step(4, "Track Score: derived columns computed from the probability rows")
def scored_on(tid, m):
    if (tid, m) == (4817, 5): return datetime(2026, 7, 1, 9, 14, 52)
    rel = T[tid][4]; dep = M[m][7]
    first = min(T[tid][7], key=lambda x: M[x][7])
    if m == first:                       # scored at ingest, before release
        d = max(dep, A[T[tid][1]][3]) + timedelta(days=1)
        d = min(d, rel - timedelta(days=1))
    else:                                # later model: backfill after its deployment
        d = max(dep, rel) + timedelta(days=rng.randint(1, 6))
    return datetime(d.year, d.month, d.day, rng.randint(0, 23), rng.randint(0, 59), rng.randint(0, 59))

TRACK_SCORE = []
for tid, *_ , mids in TRACKS:
    for m in mids:
        u = DIST[(tid, m)]
        others = [DIST[(tid, o)] for o in mids if o != m]
        div = round(sum(jsd(u, v) for v in others) / len(others), 4) if others else None
        seg_var = 0.0051 if (tid, m) == (4817, 5) else round(rng.uniform(0.0015, 0.0300), 4)
        TRACK_SCORE.append(dict(
            track=tid, model=m, genre=u.index(max(u)), top=max(u) / U,
            entropy=round(ent(u), 4), div=div, segvar=seg_var,
            on=scored_on(tid, m)))
TS = {(r["track"], r["model"]): r for r in TRACK_SCORE}
GENRE_PROB = [(tid, m, g, u[g] / U) for (tid, m), u in DIST.items() for g in range(NG)]
GENRE_PROB.sort()

# --------------------------------------------------------------------------
step(5, "stations, blends, plays")
STATIONS = [  # id, owner, model, name, min amb, max amb, created, blend
    (7710, 20931, 5, "Late Jazz Hop",        0.35, 0.80, date(2026, 8, 1),  {"jazz": .60, "hiphop": .40}),
    (7711, 20932, 4, "Porch and Pedal",      0.20, 0.70, date(2026, 6, 10), {"country": .70, "blues": .30}),
    (7712, 20932, 2, "Neon Waltz",           0.30, 0.90, date(2026, 5, 2),  {"pop": .50, "disco": .30, "classical": .20}),
    (7713, 20933, 1, "Iron Choir",           0.25, 0.85, date(2026, 5, 20), {"metal": .50, "classical": .34, "rock": .16}),
    (7714, 20934, 4, "Harbour Dub",          0.30, 0.90, date(2026, 6, 12), {"reggae": .60, "hiphop": .40}),
    (7715, 20935, 5, "Glitterball Drift",    0.20, 0.75, date(2026, 7, 20), {"disco": .34, "pop": .33, "reggae": .33}),
    (7716, 20936, 4, "Rust Belt Radio",      0.10, 0.70, date(2026, 7, 1),  {"rock": .55, "blues": .25, "country": .20}),
    (7717, 20937, 2, "Blue Hour",            0.20, 0.80, date(2026, 8, 5),  {"blues": 1.00}),
]
S = {s[0]: s for s in STATIONS}
STATION_BLEND = sorted((sid, G[g], w) for sid, *_, bl in STATIONS for g, w in bl.items())

def station_can_surface(sid, tid):
    s = S[sid]; r = TS.get((tid, s[2]))
    return (r is not None and s[4] <= r["entropy"] <= s[5]
            and GENRES[r["genre"]][1] in s[7])          # top genre is part of the station's blend

PLAYS = [  # listener, track, played_at, station(None=direct), seconds, reaction
    (20931, 4817, datetime(2026, 8, 2, 21, 47, 5),  7710, 214, "saved"),
]
REACT = [None, None, None, "saved", "skipped"]
plan = [  # (listener, station or None, n plays)
    (20931, 7710, 2), (20931, None, 2), (20932, 7711, 2), (20932, 7712, 2),
    (20933, 7713, 3), (20934, 7714, 3), (20935, 7715, 2), (20936, 7716, 2),
    (20937, 7717, 1), (20938, 7710, 2),   # 20938 uses jrivera's shared station (I6)
    (20935, 7714, 2),                      # another shared-station listener
    (20934, None, 2), (20936, None, 2), (20938, None, 2), (20933, None, 1),
]
scored_ids = [t[0] for t in TRACKS if t[5] == "scored"]
for lid, sid, n in plan:
    pool = [t for t in scored_ids if (station_can_surface(sid, t) if sid else True)]
    if not pool: sys.exit(f"station {sid} can surface no track: widen its ambiguity range")
    for _ in range(n):
        tid = rng.choice(pool)
        start = max(L[lid][4], T[tid][4], S[sid][6] if sid else date(2026, 1, 1)) + timedelta(days=rng.randint(1, 20))
        start = min(start, date(2026, 9, 23))
        at = datetime(start.year, start.month, start.day, rng.randint(0, 23), rng.randint(0, 59), rng.randint(0, 59))
        react = rng.choice(REACT)
        dur = T[tid][3]
        secs = rng.randint(8, 28) if react == "skipped" else (dur if rng.random() < 0.7 else rng.randint(40, dur - 1))
        PLAYS.append((lid, tid, at, sid, secs, react))
QUALIFY = 30  # seconds streamed for a play to earn royalty
PLAY = list(dict(listener=l, track=t, at=at, station=s, secs=sec, done=int(sec == T[t][3]),
                   react=r, royalty=A[T[t][1]][4] if sec >= QUALIFY else 0.0)
              for l, t, at, s, sec, r in PLAYS)
PLAY = sorted(PLAY, key=lambda p: (p["at"], p["listener"]))

# --------------------------------------------------------------------------
step(6, "integrity checks (keys, FKs, I1-I6, counts)")
fails = []
def check(ok, msg): 
    if not ok: fails.append(msg)
def uniq(rows, key, name): check(len({key(r) for r in rows}) == len(rows), f"duplicate PK in {name}")
uniq(ARTISTS, lambda r: r[0], "Artist"); uniq(TRACKS, lambda r: r[0], "Track")
uniq(LISTENERS, lambda r: r[0], "Listener"); uniq(STATIONS, lambda r: r[0], "Station")
uniq(TRACK_SCORE, lambda r: (r["track"], r["model"]), "Track Score")
uniq(GENRE_PROB, lambda r: r[:3], "Genre Probability")
uniq(STATION_BLEND, lambda r: r[:2], "Station Blend")
uniq(PLAY, lambda r: (r["listener"], r["track"], r["at"]), "Play")
for col, rows, i in [("Stage Name", ARTISTS, 1), ("Email", LISTENERS, 1), ("Version Label", MODELS, 1)]:
    check(len({r[i] for r in rows}) == len(rows), f"UNIQUE {col}")
check(all(t[1] in A for t in TRACKS), "FK Track.Artist")
check(all(s[1] in L and s[2] in M for s in STATIONS), "FK Station")
check(all(r["track"] in T and r["model"] in M and 0 <= r["genre"] < NG for r in TRACK_SCORE), "FK Track Score")
check(all((t, m) in TS for t, m, _, _ in GENRE_PROB), "composite FK Genre Probability -> Track Score")
check(all(p["listener"] in L and p["track"] in T and (p["station"] is None or p["station"] in S) for p in PLAY), "FK Play")
# I1 blend: >=1 row, sums to exactly 1.00
for sid in S:
    ws_ = [w for s, _, w in STATION_BLEND if s == sid]
    check(len(ws_) >= 1 and round(sum(ws_), 2) == 1.00 and all(0 < w <= 1 for w in ws_), f"I1 station {sid}")
# I2 ten rows per score, sum to 1 within 0.0010
for (t, m) in TS:
    ps = [p for tt, mm, _, p in GENRE_PROB if (tt, mm) == (t, m)]
    check(len(ps) == NG and abs(sum(ps) - 1) <= 0.0010, f"I2 score {(t, m)}")
# I3 derived values recomputed independently from the probability rows
for (t, m), r in TS.items():
    u = [round(p * U) for tt, mm, _, p in GENRE_PROB if (tt, mm) == (t, m)]
    oth = [[round(p * U) for tt, mm, _, p in GENRE_PROB if (tt, mm) == (t, o)] for (tt, o) in TS if tt == t and o != m]
    exp_div = round(sum(jsd(u, v) for v in oth) / len(oth), 4) if oth else None
    check(r["genre"] == u.index(max(u)) and abs(r["top"] - max(u) / U) < 1e-9
          and r["entropy"] == round(ent(u), 4) and r["div"] == exp_div, f"I3 score {(t, m)}")
check(TS[(4824, 4)]["div"] is None, "I3 single-model track must have empty divergence")
# I4
check(all(s[4] <= s[5] for s in STATIONS), "I4")
# I5 status vs scores
for t in TRACKS:
    check((t[5] == "scored") == any(k[0] == t[0] for k in TS), f"I5 track {t[0]}")
# I6 shared-station plays exist (documented, not constrained)
shared = [p for p in PLAY if p["station"] and S[p["station"]][1] != p["listener"]]
check(len(shared) >= 1, "I6 example missing")
# temporal / semantic sanity
for t in TRACKS: check(t[4] >= A[t[1]][3], f"release before artist joined {t[0]}")
for r in TRACK_SCORE:
    check(r["on"].date() >= M[r["model"]][7], f"scored before model deployed {(r['track'], r['model'])}")
for t in TRACKS:
    if T[t[0]][5] == "scored":
        check(min(TS[(t[0], m)]["on"].date() for m in t[7]) <= t[4], f"released before first score {t[0]}")
for s in STATIONS: check(s[6] >= L[s[1]][4] and s[6] >= M[s[2]][7], f"station dates {s[0]}")
for p in PLAY:
    check(p["at"].date() >= max(T[p["track"]][4], L[p["listener"]][4]), "play before release/signup")
    check(T[p["track"]][5] == "scored", "play of unscored track")
    check(p["secs"] <= T[p["track"]][3], "seconds played > duration")
    if p["station"]:
        check(p["at"].date() >= S[p["station"]][6], "play before station created")
        check(station_can_surface(p["station"], p["track"]), f"station {p['station']} surfaced out-of-range track {p['track']}")
    check(p["royalty"] == (A[T[p["track"]][1]][4] if p["secs"] >= QUALIFY else 0), "royalty")
    check(not (p["react"] == "skipped" and p["done"]), "skipped yet completed")
# record counts
counts = dict(Artist=len(ARTISTS), Track=len(TRACKS), Genre=len(GENRES), Listener=len(LISTENERS),
              **{"Scoring Model": len(MODELS)}, Station=len(STATIONS))
check(all(v >= 5 for v in counts.values()) and sum(counts.values()) >= 50, "record-count requirement")
# RA example values reproduced exactly
ex = TS[(4817, 5)]
check((ex["top"], ex["entropy"], ex["div"], ex["segvar"]) == (0.58, 0.4619, 0.0198, 0.0051), f"RA example drift {ex}")
if fails:
    print("FAILED:"); [print("  -", f) for f in fails]; sys.exit(1)
print(f"      all checks passed | non-associative records {sum(counts.values())} {counts}")
print(f"      Track Score {len(TRACK_SCORE)} | Genre Probability {len(GENRE_PROB)} | Station Blend {len(STATION_BLEND)} | Play {len(PLAY)} | shared-station plays {len(shared)}")

# --------------------------------------------------------------------------
step(7, "write data.xlsx in the data_template.xlsx format")
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
DATE, DT = "yyyy-mm-dd", "yyyy-mm-dd hh:mm:ss"
def sha(tid): return "0x" + hashlib.sha256(f"crossfade-audio-{tid}".encode()).hexdigest().upper()

BLOCKS = [
 ("Artist", ["Artist ID", "Stage Name", "Country", "Joined Date", "Payout Rate", "Verified"],
  [(a[0], a[1], a[2], a[3], a[4], a[5]) for a in ARTISTS], {3: DATE, 4: "0.00000"}),
 ("Track", ["Track ID", "Artist ID", "Title", "Duration", "Release Date", "Audio Checksum", "Ingest Status"],
  [(t[0], t[1], t[2], t[3], t[4], sha(t[0]), t[5]) for t in TRACKS], {4: DATE}),
 ("Genre", ["Genre ID", "Name", "Description"], GENRES, {}),
 ("Listener", ["Listener ID", "Email", "Display Name", "Plan Tier", "Signup Date"], LISTENERS, {4: DATE}),
 ("Scoring Model", ["Model ID", "Version Label", "Fusion Type", "Embedding Dimension", "Segment Length",
                    "Training Corpus", "Holdout Accuracy", "Deployed On"], MODELS, {6: "0.0000", 7: DATE}),
 ("Station", ["Station ID", "Listener ID", "Model ID", "Name", "Minimum Ambiguity", "Maximum Ambiguity", "Created On"],
  [s[:7] for s in STATIONS], {4: "0.00", 5: "0.00", 6: DATE}),
 ("Track Score", ["Track ID", "Model ID", "Genre ID", "Top Probability", "Posterior Entropy", "Model Divergence",
                  "Segment Variance", "Scored On"],
  [(r["track"], r["model"], r["genre"], r["top"], r["entropy"], r["div"], r["segvar"], r["on"]) for r in TRACK_SCORE],
  {3: "0.0000", 4: "0.0000", 5: "0.0000", 6: "0.0000", 7: DT}),
 ("Genre Probability", ["Track ID", "Model ID", "Genre ID", "Probability"], GENRE_PROB, {3: "0.0000"}),
 ("Station Blend", ["Station ID", "Genre ID", "Weight"], STATION_BLEND, {2: "0.00"}),
 ("Play", ["Listener ID", "Track ID", "Played At", "Station ID", "Seconds Played", "Completed", "Reaction", "Royalty Amount"],
  [(p["listener"], p["track"], p["at"], p["station"], p["secs"], p["done"], p["react"], p["royalty"]) for p in PLAY],
  {2: DT, 7: "0.00000"}),
]
wb = Workbook(); ws = wb.active; ws.title = "table records"
BOLD, NORM = Font(name="Calibri", size=11, bold=True), Font(name="Calibri", size=11)
widths = {}
row = 1
for name, header, rows, fmt in BLOCKS:
    # Title: bold and centred across the block's columns, as in the template,
    # using centre-across-selection instead of a merged range.
    for c in range(1, len(header) + 1):
        cell = ws.cell(row, c, name if c == 1 else None)
        cell.font = BOLD; cell.alignment = Alignment(horizontal="centerContinuous")
    row += 1
    for c, h in enumerate(header, 1):
        ws.cell(row, c, h).font = BOLD
        widths[c] = max(widths.get(c, 0), len(h))
    row += 1
    for rec in rows:
        for c, v in enumerate(rec, 1):
            cell = ws.cell(row, c, v); cell.font = NORM
            if c - 1 in fmt: cell.number_format = fmt[c - 1]
            shown = len(str(v)) if v is not None else 0
            if isinstance(v, datetime): shown = 19
            elif isinstance(v, date): shown = 10
            widths[c] = max(widths.get(c, 0), min(shown, 70))
        row += 1
    row += 1  # blank separator row
for c, w in widths.items():
    ws.column_dimensions[get_column_letter(c)].width = w + 2
import os
OUT_XLSX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "deliverables", "data.xlsx")
wb.save(OUT_XLSX)
print(f"      wrote {os.path.normpath(OUT_XLSX)}: {len(BLOCKS)} blocks, {row - 2} rows used, merged ranges: {len(ws.merged_cells.ranges)}")
